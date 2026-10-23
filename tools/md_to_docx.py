"""把笔试答卷 Markdown 转成可提交的 Word 文档。

只处理这份文档里实际出现过的语法：标题、段落、表格、代码块、引用、列表、横线。
不追求通用性，追求输出干净。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

EA_FONT = "微软雅黑"
LATIN_FONT = "Segoe UI"
CODE_FONT = "Consolas"

BOLD = re.compile(r"\*\*(.+?)\*\*")
CODE = re.compile(r"`([^`]+)`")


def set_run_font(run, name: str, size: float | None = None, bold: bool = False,
                 color: RGBColor | None = None) -> None:
    run.font.name = name
    run.font.bold = bold or None
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), EA_FONT)


def add_inline(paragraph, text: str, base_size: float = 10.5) -> None:
    """把 **粗体** 与 `代码` 拆成 run。"""
    tokens: list[tuple[str, str]] = []
    cursor = 0
    pattern = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`")
    for match in pattern.finditer(text):
        if match.start() > cursor:
            tokens.append(("plain", text[cursor:match.start()]))
        if match.group(1) is not None:
            tokens.append(("bold", match.group(1)))
        else:
            tokens.append(("code", match.group(2)))
        cursor = match.end()
    if cursor < len(text):
        tokens.append(("plain", text[cursor:]))

    for kind, value in tokens:
        if kind == "bold":
            run = paragraph.add_run(value)
            set_run_font(run, LATIN_FONT, base_size, bold=True)
        elif kind == "code":
            run = paragraph.add_run(value)
            set_run_font(run, CODE_FONT, base_size - 0.5, color=RGBColor(0xB0, 0x30, 0x60))
        else:
            run = paragraph.add_run(value)
            set_run_font(run, LATIN_FONT, base_size)


def shade(paragraph, hex_fill: str = "F4F6F8") -> None:
    ppr = paragraph._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    ppr.append(shd)


def style_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), EA_FONT)

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = document.styles[name]
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x1A, 0x2A, 0x3A)
        srpr = style.element.get_or_add_rPr()
        srfonts = srpr.find(qn("w:rFonts"))
        if srfonts is None:
            srfonts = OxmlElement("w:rFonts")
            srpr.append(srfonts)
        srfonts.set(qn("w:eastAsia"), EA_FONT)


def flush_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            value = row[c_idx] if c_idx < len(row) else ""
            add_inline(paragraph, value, base_size=9.5)
            if r_idx == 0:
                for run in paragraph.runs:
                    run.font.bold = True
    document.add_paragraph()


def build(source: Path, target: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    style_document(document)

    buffer_table: list[list[str]] = []
    in_code = False
    code_lines: list[str] = []

    def close_table() -> None:
        nonlocal buffer_table
        if buffer_table:
            flush_table(document, buffer_table)
            buffer_table = []

    def close_code() -> None:
        nonlocal code_lines
        if not code_lines:
            return
        paragraph = document.add_paragraph()
        paragraph.style = document.styles["No Spacing"]
        for index, code_line in enumerate(code_lines):
            run = paragraph.add_run(code_line)
            set_run_font(run, CODE_FONT, 9.0)
            if index < len(code_lines) - 1:
                run.add_break()
        shade(paragraph)
        code_lines = []
        document.add_paragraph()

    for line in lines:
        if line.startswith("```"):
            if in_code:
                close_code()
                in_code = False
            else:
                close_table()
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        stripped = line.strip()

        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            buffer_table.append(cells)
            continue

        close_table()

        if not stripped:
            continue

        if stripped == "---":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run("— — —")
            set_run_font(run, LATIN_FONT, 9, color=RGBColor(0x99, 0x99, 0x99))
            continue

        if stripped.startswith("#### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_inline(paragraph, stripped[5:])
            continue
        if stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_inline(paragraph, stripped[4:])
            continue
        if stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, stripped[3:])
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, stripped[2:])
            continue

        if stripped.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            run = paragraph.add_run(stripped[2:].replace("**", ""))
            set_run_font(run, LATIN_FONT, 10)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x44, 0x55, 0x66)
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(2))
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped)

    close_table()
    if in_code:
        close_code()

    document.save(target)


def verify(target: Path) -> None:
    document = Document(target)
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    print(f"段落数：{len(paragraphs)}")
    print(f"表格数：{len(document.tables)}")
    print("前几段：")
    for paragraph in paragraphs[:6]:
        print(f"  [{paragraph.style.name}] {paragraph.text[:50]}")
    assert len(document.tables) >= 6, f"表格数量异常（{len(document.tables)}），转换可能丢了内容"
    assert any("单选题" in p.text for p in paragraphs), "缺少正文标题"
    code_like = [p for p in paragraphs if p.style.name == "No Spacing"]
    assert code_like, "代码块没有转换出来"


if __name__ == "__main__":
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    build(source, target)
    verify(target)
    print(f"已生成：{target}")
